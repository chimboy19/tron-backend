# backend/generate_sample_data_extended.py
import os
import random
from datetime import datetime, timedelta
from openpyxl import Workbook
import pandas as pd # Using pandas for easier handling if needed, but openpyxl is sufficient here

# Ensure data directory exists
os.makedirs("data", exist_ok=True)

def save_xlsx(filepath, headers, rows):
    """Generic function to save data to an XLSX file."""
    wb = Workbook()
    ws = wb.active
    ws.title = os.path.basename(filepath).split('.')[0] # Use filename stem as sheet name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(filepath)
    print(f"✅ Created {filepath}")

def generate_sample_raf1():
    """Generate sample data for RAF1.xlsx (FAX History)."""
    filepath = "data/RAF1.xlsx"
    headers = [
        "RAF1D", "FAXDAY", "FAXTIM", "JOBNBR", "SPLNO", "LOGSPL", "MCOD", "ODRNO",
        "FAXSU", "FLAG", "FLAGS", "SNDNO", "FLAGP", "FLAG1", "RAF1XA", "RAF1XB",
        "RAF1XC", "RAF1XD", "RAF1XE", "JOBNBA", "SPLNOA", "JOB", "USER", "UFSNO",
        "UCOD", "UFNO", "GPCOD", "WCODF", "UFMUC"
    ]

    # Sample MCODs and UCODs to link to existing data (or create placeholders)
    # These should ideally match MCODs in MMK1.xlsx and UCODs in MUS1.xlsx
    sample_mcods = [f"SUP{i:02d}" for i in range(1, 6)] # e.g., SUP01, SUP02, ...
    sample_ucods = [f"C{i:04d}" for i in range(1001, 1011)] # e.g., C1001, C1002, ...

    rows = []
    for i in range(20): # Generate 20 sample records
        date_obj = datetime.now() - timedelta(days=random.randint(1, 30))
        time_obj = datetime.now().replace(hour=random.randint(8, 18), minute=random.randint(0, 59), second=random.randint(0, 59))
        rows.append([
            "", # RAF1D (Delete Flag) - Usually empty or 'D'
            date_obj.strftime("%Y%m%d"), # FAXDAY (YYYYMMDD)
            time_obj.strftime("%H%M%S"), # FAXTIM (HHMMSS)
            f"JOB{i:06d}", # JOBNBR
            f"SPL{i:04d}", # SPLNO (Spool Number Print)
            f"LOG{i:04d}", # LOGSPL (Spool Number Log)
            random.choice(sample_mcods), # MCOD (Supplier Code)
            f"ORD{i:05d}", # ODRNO (Order Number Sent To)
            random.randint(0, 3), # FAXSU (Resend Count)
            random.choice(["", "P"]), # FLAG (Printer Output Flag)
            random.choice(["", "Y"]), # FLAGS (Resend Flag)
            f"SN{i:06d}", # SNDNO (Send Number)
            random.choice(["", "P"]), # FLAGP (Printer Flag P)
            random.choice(["", "1"]), # FLAG1 (FAX Type Flag)
            random.choice(["", "X", "Y"]), # RAF1XA
            random.choice(["", "X", "Y"]), # RAF1XB
            random.choice(["", "X", "Y"]), # RAF1XC
            random.choice(["", "X", "Y"]), # RAF1XD
            random.choice(["", "1", "2"]), # RAF1XE
            f"JB{i:06d}", # JOBNBA (Job Number Inquiry)
            f"SL{i:04d}", # SPLNOA (Spool Number Sent)
            f"J{i:03d}", # JOB (Job)
            f"U{i:03d}", # USER (User)
            f"UFS{i:07d}", # UFSNO (FAX Management Number)
            random.choice(sample_ucods), # UCOD (Customer Code)
            f"UF{i:04d}", # UFNO (Destination Code)
            f"G{i:03d}", # GPCOD (Sending Department)
            f"W{i:03d}", # WCODF (Sender Code)
            random.choice(["", "S", "C"]), # UFMUC (Destination Type: S=Supplier, C=Customer)
        ])

    save_xlsx(filepath, headers, rows)

def generate_sample_tnz2():
    """Generate sample data for TNZ2.xlsx (Customer Delivery Schedule)."""
    filepath = "data/TNZ2.xlsx"
    headers = [
        "TNZ2D", "DENNO", "GYONO", "UCOD", "MCOD", "HNAME", "SURYO", "TANKA", "IRINE",
        "NODAYU", "SYDAY", "UNMX", "HNM2", "MNO", "KNO", "MNO2", "KNO2", "BURIK", "BGPSY", "SOKO",
        "ZLVL", "TANAFL", "TANABL", "TANANO", "TANAST", "HAISO", "KTANA"
    ]

    # Sample MCODs and UCODs to link to existing data (or create placeholders)
    sample_mcods = [f"SUP{i:02d}" for i in range(1, 6)]
    sample_ucods = [f"C{i:04d}" for i in range(1001, 1011)]
    # Sample HCODs to link to MHN1 (or create placeholders)
    sample_hcods = [f"H{i:06d}" for i in range(100000, 100020)] # e.g., H100000, H100001, ...

    rows = []
    for i in range(25): # Generate 25 sample records
        date_obj = datetime.now() + timedelta(days=random.randint(5, 60)) # Future delivery date
        shipment_date_obj = date_obj - timedelta(days=random.randint(1, 5)) # Shipment date before delivery
        rows.append([
            "", # TNZ2D (Delete Flag) - Usually empty or 'D'
            f"D{i:06d}", # DENNO (Document Number)
            f"G{i:02d}", # GYONO (Line Number)
            random.choice(sample_ucods), # UCOD (Customer Code)
            random.choice(sample_mcods), # MCOD (Supplier Code)
            f"Item {random.choice(sample_hcods)}", # HNAME (Item Description/Note)
            random.randint(1, 100), # SURYO (Quantity)
            round(random.uniform(100.0, 10000.0), 2), # TANKA (Unit Price)
            round(random.uniform(80.0, 8000.0), 2), # IRINE (Standard Cost)
            date_obj.strftime("%Y%m%d"), # NODAYU (Customer Delivery Date YYYYMMDD)
            shipment_date_obj.strftime("%Y%m%d"), # SYDAY (Shipment Date YYYYMMDD)
            f"Customer {random.choice(sample_ucods)}", # UNMX (Customer Name for Various Ports)
            f"Model {random.choice(sample_hcods)[-3:]}", # HNM2 (Model 2)
            f"M{i:06d}", # MNO (Estimate Number)
            f"K{i:05d}", # KNO (Job Number)
            f"M{i:06d}-{random.randint(1, 10)}", # MNO2 (Estimate Number Detail)
            f"K{i:05d}-{random.randint(1, 10)}", # KNO2 (Job Number Detail)
            random.choice(["A", "B", "C"]), # BURIK (Sales Category)
            random.choice(["", "Y"]), # BGPSY (Group Identifier: Y=Yes?)
            random.choice(["01", "02", "03"]), # SOKO (Warehouse Code)
            random.choice(["1", "2", "3"]), # ZLVL (Inventory Level)
            random.choice(["A", "B", "C"]), # TANAFL (Location Floor)
            random.choice(["1", "2"]), # TANABL (Location Block)
            f"{random.randint(1, 99):02d}", # TANANO (Location Number)
            f"{random.randint(1, 9):01d}", # TANAST (Location Stage)
            random.choice(["DEL", "SHIP", "PICK"]), # HAISO (Delivery Code)
            f"KTA{i:06d}", # KTANA (Storage Shelf)
        ])

    save_xlsx(filepath, headers, rows)

if __name__ == "__main__":
    generate_sample_raf1()
    generate_sample_tnz2()
    print("\nSample RAF1.xlsx and TNZ2.xlsx files generated in the 'data' directory.")
    print("Ensure the field names in the XLSX match the keys in the load_xlsx_to_db function calls in app.py/utils/xlsx_loader.py.")
    print("You may need to adjust the sample data values to match your actual supplier/customer/product codes.")